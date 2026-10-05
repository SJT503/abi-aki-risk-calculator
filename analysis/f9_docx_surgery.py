# -*- coding: utf-8 -*-
"""Round-9.1 docx surgery: apply the md-level patch to the delivered docx directly.
Run-level replacement preserves character formatting outside the edited span;
table row insertion clones the anchor row's XML so borders/widths carry over.
Files: MANUSCRIPT_npjDM_round9.docx, TABLES_npjDM_round9.docx, Supplementary_Information_round9.docx
"""
import copy
import docx

BASE = "E:/TBI subtype/09_tbi_aki/_biomni_r9/round9_revision"

def replace_in_runs(para, old, new):
    """Replace `old` -> `new` inside one paragraph, preserving runs outside the span."""
    full = "".join(r.text for r in para.runs)
    if old not in full:
        return False
    s = full.index(old)
    e = s + len(old)
    pos = 0
    first = True
    for r in para.runs:
        rs, re_ = pos, pos + len(r.text)
        if re_ <= s or rs >= e:
            pos = re_
            continue
        local_s, local_e = max(s - rs, 0), min(e - rs, len(r.text))
        keep_pre = r.text[:local_s]
        keep_post = r.text[local_e:]
        r.text = (keep_pre + new + keep_post) if first else (keep_pre + keep_post)
        first = False
        pos = re_
    return True

def replace_in_doc(doc, old, new, expect=1):
    n = 0
    for p in doc.paragraphs:
        if old in p.text:
            n += replace_in_runs(p, old, new)
    assert n == expect, f"expected {expect} replacements, made {n}: {old[:60]!r}"

def fix_table_row(doc, row_vals, new_vals):
    """Find the table row whose cells equal row_vals; rewrite to new_vals."""
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if cells == row_vals:
                for c, v in zip(row.cells, new_vals):
                    c.text = v
                return row
    raise AssertionError(f"row not found: {row_vals}")

def insert_row_after(doc, anchor_vals, new_vals):
    for t in doc.tables:
        for row in t.rows:
            if [c.text.strip() for c in row.cells] == new_vals[:0] + [c.text.strip() for c in row.cells]:
                pass
        for i, row in enumerate(t.rows):
            if [c.text.strip() for c in row.cells][:5] == anchor_vals[:5] and len(anchor_vals) <= len(row.cells):
                if [c.text.strip() for c in row.cells] == anchor_vals:
                    new_tr = copy.deepcopy(row._tr)
                    row._tr.addnext(new_tr)
                    new_row = t.rows[i + 1]
                    for c, v in zip(new_row.cells, new_vals):
                        c.text = v
                    return new_row
    raise AssertionError(f"anchor row not found: {anchor_vals}")

ROW_OLD = ["Tier 2 − Tier 1", "MIMIC-IV internal test", "not estimated", "not estimated", "not estimated"]
ROW_NEW = ["Tier 2 − Tier 1", "MIMIC-IV internal test", "+0.064", "−0.151 to 0.198", "0.548"]
ROW2_NEW = ["Tier 2 − ≥Stage 2", "MIMIC-IV internal test", "+0.069", "−0.123 to 0.251", "0.516"]
FN_OLD = "The raw internal Tier 2 − Tier 1 difference was not estimated in the current analysis files."
FN_NEW = "Internal paired differences use the same resampling scheme (analysis file f9_fillins.json, patch 9.1)."
RES_OLD = ("on 40 positive checkpoints from 15 event stays; after proximity standardisation the internal "
           "difference was +0.024 (−0.164 to 0.160; P = 0.716), compatible with both an advantage and none.")
RES_NEW = ("on 40 positive checkpoints from 15 event stays; the raw paired internal difference was +0.064 "
           "(−0.151 to 0.198; P = 0.548), directionally consistent with the external advantage but imprecise, "
           "and after proximity standardisation the internal difference was +0.024 (−0.164 to 0.160; P = 0.716), "
           "compatible with both an advantage and none.")
LIM_OLD = ("(proximity-standardised difference +0.024, −0.164 to 0.160; P = 0.716), and the raw internal "
           "paired difference was not estimated.")
LIM_NEW = ("(raw paired difference +0.064, −0.151 to 0.198, P = 0.548; proximity-standardised difference "
           "+0.024, −0.164 to 0.160, P = 0.716).")
DIS_OLD = ("Calibrated probabilities and deployment views were usable for Tier 1. The route from here is "
           "concrete: replication of the severe tier in a second system, calibration and threshold setting "
           "for the Tier-2 output, silent-mode operation alongside the existing record system, and a "
           "stepped-wedge evaluation")
DIS_NEW = ("Calibrated probabilities and deployment views were usable for Tier 1, and a matching "
           "development-fitted recalibration layer gave the Tier-2 output a usable scale (external slope "
           "1.287, expected calibration error 0.0022), at which a 0.05 severe-tier threshold captured 43.4% "
           "of external severe events a median 6 h ahead while alerting 2.0% of stays (Supplementary "
           "Material 5). The route from here is concrete: replication of the severe tier in a second system, "
           "prospective confirmation of Tier-2 thresholds, silent-mode operation alongside the existing "
           "record system, and a stepped-wedge evaluation")
SM5FN_OLD = "The raw internal paired differences involving ≥Stage 3 were not estimated in the current analysis files; the internal severe-tier comparison is reported after proximity standardisation (below)."
SM5FN_NEW = ("Raw internal paired differences involving ≥Stage 3 were added in patch 9.1 (analysis file "
             "f9_fillins.json): Tier 2 − Tier 1 = +0.064 (−0.151 to 0.198; P = 0.548) and Tier 2 − ≥Stage 2 = "
             "+0.069 (−0.123 to 0.251; P = 0.516), directionally consistent with the external advantage but "
             "imprecise at 15 severe-event stays; the proximity-standardised comparison is reported below.")

# ================= MANUSCRIPT =================
f = f"{BASE}/manuscript/MANUSCRIPT_npjDM_round9.docx"
doc = docx.Document(f)
fix_table_row(doc, ROW_OLD, ROW_NEW)
anchor_after = ["Tier 2 − Tier 1", "MIMIC-IV internal test", "+0.064", "−0.151 to 0.198", "0.548"]
insert_row_after(doc, anchor_after, ROW2_NEW)
replace_in_doc(doc, FN_OLD, FN_NEW)
replace_in_doc(doc, RES_OLD, RES_NEW)
replace_in_doc(doc, LIM_OLD, LIM_NEW)
replace_in_doc(doc, DIS_OLD, DIS_NEW)
doc.save(f)
print("MANUSCRIPT docx patched")

# ================= TABLES =================
f = f"{BASE}/manuscript/TABLES_npjDM_round9.docx"
doc = docx.Document(f)
fix_table_row(doc, ROW_OLD, ROW_NEW)
insert_row_after(doc, anchor_after, ROW2_NEW)
replace_in_doc(doc, FN_OLD, FN_NEW)
doc.save(f)
print("TABLES docx patched")

# ================= SI (merged) =================
f = f"{BASE}/supplementary/Supplementary_Information_round9.docx"
doc = docx.Document(f)
replace_in_doc(doc, SM5FN_OLD, SM5FN_NEW)

# insert the new Tier-2 section before the 'Comparison with an admission-variable baseline' heading
target = None
for p in doc.paragraphs:
    if p.text.strip() == "Comparison with an admission-variable baseline":
        target = p
        break
assert target is not None, "anchor heading not found"

def add_para_before(anchor, text, style=None, bold=False, italic=False):
    new_p = anchor.insert_paragraph_before(text, style=style)
    if bold:
        for r in new_p.runs: r.bold = True
    if italic:
        for r in new_p.runs: r.italic = True
    return new_p

# find a body/table style used in the SI for headings and tables
heading_style = None
for p in doc.paragraphs:
    if p.text.strip().startswith("Comparison with an admission-variable baseline"):
        heading_style = p.style
        break
sec_paras = [
    ("Tier-2 (severe AKI) output: recalibration layer and threshold views", heading_style, False),
    ("Because the Tier-2 model shares the class-weighted training of Tier 1, its raw probabilities are over-dispersed relative to the 0.3% checkpoint event rate. A recalibration layer was fitted with the identical protocol used for Tier 1 (logistic recalibration on out-of-fold predictions from five-fold GroupKFold by patient on the training split; analysis file f9_tier2_calibration.json, patch 9.1): intercept −0.301766, slope 0.393865. Applied frozen to the internal test set and to eICU-CRD:", None, False),
]
# insert_paragraph_before stacks in order when always inserting before the SAME anchor
for text, style, _ in sec_paras:
    add_para_before(target, text, style=style)

# tables: build after the intro paragraphs -> insert before anchor as paragraphs can't hold tables directly;
# python-docx: create table at end then move its XML before anchor.
def add_table_before(doc, anchor_para, rows, header):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    try:
        t.style = doc.tables[0].style
    except Exception:
        pass
    for j, h in enumerate(header):
        t.rows[0].cells[j].text = h
        for r in t.rows[0].cells[j].paragraphs[0].runs: r.bold = True
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v
    anchor_para._p.addprevious(t._tbl)
    return t

T1_HEADER = ["Cohort", "Slope before", "Slope after", "ECE before", "ECE after"]
T1_ROWS = [["MIMIC-IV internal test", "0.429", "1.067", "0.0015", "0.0012"],
           ["eICU-CRD external", "0.506", "1.287", "0.0030", "0.0022"]]
add_table_before(doc, target, T1_ROWS, T1_HEADER)
add_para_before(target, "The slight over-correction externally (slope above 1) mirrors Tier 1 (1.159). Threshold views below use the deployment definition version 2 (Methods), anchored on each event stay's final Tier-2-positive checkpoint; event-free patient-days follow the documented denominator.")
T2_HEADER = ["Threshold", "Capture", "Lead, median (IQR) h", "NNE", "False alerts /100 event-free patient-days", "Stays alerted"]
T2_ROWS = [["0.05", "0.434", "6.0 (6.0–18.0)", "4.97", "0.46", "2.0%"],
           ["0.10", "0.253", "12.0 (6.0–24.0)", "5.52", "0.30", "1.3%"],
           ["0.15", "0.193", "12.0 (6.0–25.5)", "4.25", "0.16", "0.7%"],
           ["0.20", "0.133", "12.0 (6.0–18.0)", "4.09", "0.10", "0.5%"]]
add_table_before(doc, target, T2_ROWS, T2_HEADER)
add_para_before(target, "A severe-tier threshold near 0.05 offers the best trade-off in this external view: 43.4% of severe events captured with alerts on 2.0% of stays. These are development-side views for planning silent-mode operation, not prospectively validated action thresholds.")
doc.save(f)
print("SI docx patched")
