# -*- coding: utf-8 -*-
"""Round-9.2 docx surgery (retry): md already patched; asterisks are formatting in docx,
so footnote anchors drop the md asterisk. All other anchors unchanged.
"""
from pathlib import Path
import docx

FIN = Path("E:/TBI subtype/09_tbi_aki/_biomni_r9_1/round9_1_review/round9_revision_91_final")

def replace_in_runs(para, old, new):
    full = "".join(r.text for r in para.runs)
    if old not in full:
        return False
    s, e = full.index(old), full.index(old) + len(old)
    pos, first = 0, True
    for r in para.runs:
        rs, re_ = pos, pos + len(r.text)
        if re_ <= s or rs >= e:
            pos = re_
            continue
        ls, le = max(s - rs, 0), min(e - rs, len(r.text))
        r.text = (r.text[:ls] + new + r.text[le:]) if first else (r.text[:ls] + r.text[le:])
        first = False
        pos = re_
    return True

def replace_in_doc(doc, old, new, expect=1):
    n = sum(replace_in_runs(p, old, new) for p in doc.paragraphs)
    assert n == expect, f"expected {expect} got {n}: {old[:60]!r}"

MANU = [
 ("; the Tier-2 output was not recalibrated.",
  "; the Tier-2 output received an identically fitted recalibration layer (Supplementary Material 5)."),
 ("The second step is a proposal for prospective testing: the Tier-2 output is not recalibrated, its operating threshold has not been set, and its evidence comes from one external database.",
  "The second step is a proposal for prospective testing: the Tier-2 output now carries a development-fitted recalibration layer with development-side threshold views (Supplementary Material 5), but its action threshold is not prospectively validated and its evidence comes from one external database."),
 ("Tier-2 probabilities were not recalibrated.",
  "Tier-2 probabilities received an identically fitted layer (a = −0.301766, b = 0.393865; Supplementary Material 5)."),
 ("The Tier-2 output was not recalibrated.",
  "The Tier-2 output received an identically fitted layer (Supplementary Material 5)."),
 ("Internal paired differences use the same resampling scheme (analysis file f9_fillins.json, patch 9.1).",
  "Internal paired differences use the same resampling scheme as the external comparisons."),
 ("at which a 0.05 severe-tier threshold captured 43.4% of external severe events a median 6 h ahead while alerting 2.0% of stays (Supplementary Material 5)",
  "at which the lowest examined severe-tier threshold (0.05) captured 36 of 83 external severe-event stays (43.4%) a median 6 h—one checkpoint—ahead while alerting 179 of 9,071 stays (2.0%) (Supplementary Material 5)"),
]
f = FIN / "manuscript/MANUSCRIPT_npjDM_round9.docx"
doc = docx.Document(f)
for old, new in MANU:
    replace_in_doc(doc, old, new)
doc.save(f)
print("MANUSCRIPT docx patched (6)")

TAB = [
 ("Internal paired differences use the same resampling scheme (analysis file f9_fillins.json, patch 9.1).",
  "Internal paired differences use the same resampling scheme as the external comparisons."),
 ("The Tier-2 output was not recalibrated.",
  "The Tier-2 output received an identically fitted layer (Supplementary Material 5)."),
]
f = FIN / "manuscript/TABLES_npjDM_round9.docx"
doc = docx.Document(f)
for old, new in TAB:
    replace_in_doc(doc, old, new)
doc.save(f)
print("TABLES docx patched (2)")

f = FIN / "supplementary/Supplementary_Information_round9.docx"
doc = docx.Document(f)
replace_in_doc(doc,
  "Raw internal paired differences involving ≥Stage 3 were added in patch 9.1 (analysis file f9_fillins.json): Tier 2 − Tier 1",
  "Raw internal paired differences involving ≥Stage 3, estimated with the same resampling scheme as the external comparisons: Tier 2 − Tier 1")
replace_in_doc(doc,
  "(logistic recalibration on out-of-fold predictions from five-fold GroupKFold by patient on the training split; analysis file f9_tier2_calibration.json, patch 9.1): intercept",
  "(logistic recalibration on out-of-fold predictions from five-fold GroupKFold by patient on the training split, identical to the Tier-1 protocol): intercept")
replace_in_doc(doc,
  "A severe-tier threshold near 0.05 offers the best trade-off in this external view: 43.4% of severe events captured with alerts on 2.0% of stays.",
  "The 0.05 threshold was the lowest examined and captured the largest share of severe event stays (36/83, 43.4%) with alerts on 179 of 9,071 stays (2.0%); its median lead of 6 h equals the checkpoint spacing, so most captures preceded the anchor by a single six-hour checkpoint.")

def replace_cell(doc, old, new, expect=1):
    n = 0
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                if c.text.strip() == old:
                    c.text = new
                    n += 1
    assert n == expect, f"cell expected {expect} got {n}: {old!r}"

for old, new in [("0.434", "0.434 (36/83)"), ("0.253", "0.253 (21/83)"), ("0.193", "0.193 (16/83)"), ("0.133", "0.133 (11/83)"),
                 ("2.0%", "2.0% (179/9,071)"), ("1.3%", "1.3% (116/9,071)"), ("0.7%", "0.7% (68/9,071)"), ("0.5%", "0.5% (45/9,071)")]:
    replace_cell(doc, old, new)

n_sm4 = 0
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            txt = c.text.strip()
            if txt == "MS Results, Calibration (slope/ECE/Brier); Table 3; Fig. 3a, b":
                c.text = "MS Results, Calibration (slope/ECE/Brier for Tier 1); Table 3; Fig. 3a, b; Tier-2 layer in Supplementary Material 5"
                n_sm4 += 1
            elif txt == "MS Methods, Calibration (recalibration layer fitted on patient-grouped out-of-fold development predictions; shipped frozen)":
                c.text = "MS Methods, Calibration (recalibration layers for both tiers fitted on patient-grouped out-of-fold development predictions; shipped frozen; Tier-2 layer in Supplementary Material 5)"
                n_sm4 += 1
assert n_sm4 == 2, f"SM4 cells patched: {n_sm4}"
doc.save(f)
print("SI docx patched (3 paras + 8 cells + 2 SM4 cells)")
print("DOCX SURGERY 9.2 COMPLETE")
