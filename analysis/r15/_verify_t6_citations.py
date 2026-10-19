#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""INDEPENDENT VERIFICATION T6: citation order, CN/EN number alignment,
cross-file consistency for the new v5.3 content."""
import re
import sys
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')
MD = r'E:/TBI subtype/09_tbi_aki/_rewrite_r12/MANUSCRIPT_npjDM_v5.md'
DX = r'E:/TBI subtype/09_tbi_aki/_rewrite_r12/MANUSCRIPT_v5_中文审阅版.docx'
md = open(MD, encoding='utf-8').read()

print('='*80); print('T6a  citation first-appearance order'); print('='*80)
body = md.split('## References')[0]
cites = []
for m in re.finditer(r'\[(\d+(?:[–,\s]+\d+)*)\]', body):
    for part in re.split(r'[,\s]+', m.group(1)):
        if not part:
            continue
        if '–' in part:
            a, b = part.split('–')
            cites.extend(range(int(a), int(b) + 1))
        else:
            cites.append(int(part))
first = []
for c in cites:
    if c not in first:
        first.append(c)
print('first-appearance sequence:', first)
ok = first == sorted(first)
print('ascending (no forward/citation-order violation):', 'MATCH' if ok else '*** VIOLATION ***')
ref_nums = [int(m) for m in re.findall(r'^(\d+)\.', md.split('## References')[1], re.M)]
print('reference list numbers:', ref_nums[:5], '...', ref_nums[-3:],
      'complete 1..32:', ref_nums == list(range(1, 33)))
print('any citation > 32 or < 1:', [c for c in cites if c < 1 or c > 32])

print(); print('='*80); print('T6b  Chinese docx number alignment (new content)'); print('='*80)
doc = Document(DX)
joined = '\n'.join(p.text for p in doc.paragraphs)
# probe numbers from the new v5.3 content: (label, list of acceptable strings)
probes = [
    ('captured MIMIC 239', ['239']),
    ('captured eICU 260', ['260']),
    ('missed 492', ['492']),
    ('missed 487', ['487']),
    ('eventfree 6711', ['6,711', '6711']),
    ('eventfree 8324', ['8,324', '8324']),
    ('mort MIMIC 38.9/21.7/9.8', ['38.9', '21.7', '9.8']),
    ('mort eICU 27.7/23.6/6.0', ['27.7', '23.6', '6.0']),
    ('LOS 6.47/5.45/3.40', ['6.47', '5.45', '3.40']),
    ('LOS 6.91/5.21/2.89', ['6.91', '5.21', '2.89']),
    ('quartiles MIMIC 2.9/6.1/12.9/24.3', ['2.9', '6.1', '12.9', '24.3']),
    ('quartiles eICU 2.0/4.6/8.2/15.5', ['2.0', '4.6', '8.2', '15.5']),
    ('comparator 0.4%', ['0.4']),
    ('comparator 0%', ['0%']),
    ('SMD 0.10', ['0.10', '0.1']),
    ('median 18 h', ['18']),
    ('capture pct 32.7/34.8', ['32.7', '34.8']),
    ('AUROC 0.739', ['0.739']),
    ('AUROC 0.752', ['0.752']),
    ('ablation 0.675/0.672', ['0.675', '0.672']),
    ('ablation severe 0.824/0.825', ['0.824', '0.825']),
    ('archive sha or 复现', ['复现', 'archiv', '归档']),
]
bad = []
for label, alts in probes:
    hit = any(a in joined for a in alts)
    print(f'{label:42s}', 'OK' if hit else 'MISSING')
    if not hit:
        bad.append(label)
print('missing probes:', bad if bad else 'none')

print(); print('='*80); print('T6c  EN-CN spot: clinical-outcome paragraph in docx'); print('='*80)
for p in doc.paragraphs:
    if '38.9' in p.text or '2.9' in p.text and '24.3' in p.text:
        print('>>', p.text[:260]); print()

print('='*80); print('T6d  abstract word-count detail (T5d transparency)'); print('='*80)
abst = re.search(r'## Abstract\n\n(.*?)\n\n## ', md, re.S).group(1)
toks = abst.split()
letter_words = [t for t in toks if re.search(r'[A-Za-z]', t)]
nonletter = [t for t in toks if not re.search(r'[A-Za-z]', t)]
print('letter-words %d | non-letter tokens %d | total whitespace tokens %d'
      % (len(letter_words), len(nonletter), len(toks)))
print('non-letter tokens:', nonletter)

print(); print('='*80); print('T6e  ms JSON-value spot checks'); print('='*80)
# comparator claim in ms + SI note
print('ms "0.4% and 0%" comparator:', '0.4% and 0%' in md or ('0.4%' in md and '0%' in md))
# SI note says same
si = open(r'E:/TBI subtype/09_tbi_aki/_rewrite_r12/Supplementary_Information_v5.md', encoding='utf-8').read()
print('SI note comparator 0.4% / 0%:', '0.4% (MIMIC-IV)' in si and '0% (eICU-CRD)' in si)
# SMD wording vs max SMD 0.0995
print('ms SMD sentence "at or below 0.10":', 'at or below 0.10' in md)
# lead-time wording
print('ms IQR 6-48 / 6-36:', '6–48' in md and '6–36' in md)
# model archive sentence (EN) vs docx probe
print('EN archive sentence ok:', 'model objects for all three endpoints' in md)
